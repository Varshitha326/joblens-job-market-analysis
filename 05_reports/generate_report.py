import os
import pandas as pd
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PInches

df = pd.read_csv("data/cleaned/jobs_clean.csv")
models = pd.read_csv("reports/model_comparison.csv")

# ---- Numbers computed from the data (nothing typed by hand) ----
skills = df["required_skills"].str.split("|").explode().str.strip().value_counts()
order = ["Entry (0-2 yrs)", "Mid (3-5 yrs)", "Senior (6-9 yrs)", "Lead (10+ yrs)"]
med_exp = df.groupby("experience_level")["annual_salary_usd"].median().reindex(order)
flagged = int(df["salary_outside_range"].sum())
top_countries = df["country"].value_counts().head(3)
best = models.sort_values("Macro F1", ascending=False).iloc[0]

findings = [
    f"Postings analysed: {len(df)} across {df['country'].nunique()} countries "
    f"and {df['job_category'].nunique()} job categories.",
    f"Most required skill: {skills.index[0]} ({skills.iloc[0]} postings), "
    f"followed by {skills.index[1]} and {skills.index[2]}.",
    f"Median salary rises from ${med_exp.iloc[0]:,.0f} (entry) to "
    f"${med_exp.iloc[-1]:,.0f} (lead).",
    f"Countries with most postings: " +
    ", ".join(f"{k} ({v})" for k, v in top_countries.items()) + ".",
    f"Data quality: {flagged} rows ({flagged / len(df):.0%}) have a salary outside "
    f"their own min-max range; they were flagged, not removed.",
    f"Best model by Macro F1: {best['Model']} (accuracy {best['Accuracy']}, "
    f"Macro F1 {best['Macro F1']}).",
]

charts = [
    ("Most Required Skills", "reports/top_skills.png"),
    ("Postings by Country", "reports/postings_by_country.png"),
    ("Median Salary by Experience Level", "reports/salary_by_experience.png"),
    ("Median Salary by Job Category", "reports/salary_by_category.png"),
    ("Top Skills in Entry-Level Postings", "reports/entry_level_skills.png"),
]
charts = [(t, p) for t, p in charts if os.path.exists(p)]

# ---- Word report ----
doc = Document()
doc.add_heading("JobLens: AI/ML Job Market Analysis Report", 0)
doc.add_heading("Key Findings", 1)
for f in findings:
    doc.add_paragraph(f, style="List Bullet")

doc.add_heading("Charts", 1)
for title, path in charts:
    doc.add_heading(title, 2)
    doc.add_picture(path, width=Inches(5.8))

doc.add_heading("Model Comparison", 1)
doc.add_paragraph("Task: predict job category from required skills. The baseline "
                  "always predicts the most common category.")
t = doc.add_table(rows=1, cols=len(models.columns))
t.style = "Light Grid Accent 1"
for i, c in enumerate(models.columns):
    t.rows[0].cells[i].text = c
for _, r in models.iterrows():
    cells = t.add_row().cells
    for i, v in enumerate(r):
        cells[i].text = f"{v:.3f}" if isinstance(v, float) else str(v)
doc.save("reports/JobLens_Report.docx")

# ---- PowerPoint deck ----
prs = Presentation()

s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "JobLens"
s.placeholders[1].text = f"AI/ML job market analysis of {len(df)} postings"

s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Key Findings"
tf = s.placeholders[1].text_frame
tf.text = findings[0]
for f in findings[1:]:
    tf.add_paragraph().text = f
for p in tf.paragraphs:
    for run in p.runs:
        run.font.size = PInches(0.22)

for title, path in charts:
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = title
    pic = s.shapes.add_picture(path, 0, PInches(1.5), height=PInches(5.4))
    pic.left = int((prs.slide_width - pic.width) / 2)

s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = "Model Comparison"
rows, cols = len(models) + 1, len(models.columns)
tbl = s.shapes.add_table(rows, cols, PInches(0.5), PInches(1.8),
                         PInches(9), PInches(0.5 * rows)).table
for i, c in enumerate(models.columns):
    tbl.cell(0, i).text = c
for ri, (_, r) in enumerate(models.iterrows(), start=1):
    for ci, v in enumerate(r):
        tbl.cell(ri, ci).text = f"{v:.3f}" if isinstance(v, float) else str(v)

prs.save("reports/JobLens_Deck.pptx")
print("Reports generated: reports/JobLens_Report.docx and reports/JobLens_Deck.pptx")